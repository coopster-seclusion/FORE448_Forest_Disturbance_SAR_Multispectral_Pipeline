"""s17 -- import label_tool CSV exports into the block workbooks (counts) and keep dot-level labels for the ML stretch.

Usage: python s17e_import_labels.py <export.csv> [interp2]
Writes block_status, canopy_before, lost_after, cause, confidence, offset_seen, notes into V3_blocks_labelling.xlsx
(or V3_blocks_interp2.xlsx) for every block in the export; blocks not in the export are left as they are.
Dot-level labels go to sample_blocks/dot_labels_<interp>.csv.
"""
import sys
import pandas as pd
from openpyxl import load_workbook
from v3cfg import V3

src, which = sys.argv[1], (sys.argv[2] if len(sys.argv) > 2 else "interp1")
book = f"{V3}/sample_blocks/" + ("V3_blocks_interp2.xlsx" if which == "interp2" else "V3_blocks_labelling.xlsx")
ex = pd.read_csv(src, dtype={"notes": str}).fillna("")
wb = load_workbook(book); ws = wb["Labels"]
col = {c.value: c.column for c in ws[1]}
rows = {ws.cell(r, 1).value: r for r in range(2, ws.max_row + 1)}
n = 0
for x in ex.itertuples():
    r = rows[x.block_id]
    for k, v in (("block_status", x.block_status), ("canopy_before", x.canopy_before), ("lost_after", x.lost_after),
                 ("cause", x.cause), ("confidence", x.confidence), ("offset_seen", x.offset_seen), ("notes", x.notes)):
        ws.cell(r, col[k], (int(v) if k in ("canopy_before", "lost_after") and str(v) != "" else (v or None)))
    n += 1
wb.save(book)
dots = ex[["block_id", "dots"]].assign(dot=lambda d: d.dots.str.split("|")).explode("dot")
dots["dot_no"] = dots.groupby("block_id").cumcount() + 1
dots[["block_id", "dot_no", "dot"]].to_csv(f"{V3}/sample_blocks/dot_labels_{which}.csv", index=False)
print(f"imported {n} blocks into {book}; dot labels: {len(dots)} rows")
