"""s16 -- blind pixel-level re-check of the reference points that drive the location sensitivity (log 6h).

Why: 64% of the strict estimate comes from six "Loss" points in map no-loss strata whose own 10 m pixel
shows no NDVI drop, while a neighbouring pixel does. A rule written after seeing the data (one-pixel
tolerance, both directions) moves 20 points and the estimate from 1,037 to 506 ha. Instead of adopting that
rule, the 20 points are re-labelled blind against a stricter, pre-stated pixel rule.

Design (fixed before re-labelling):
- Points = the 20 points the tolerance rule reclassifies (both directions) + 10 decoys drawn at random
  (seed 20230227) from the other Loss / No loss points, shuffled and renamed R-01..R-30.
- Chips show a 200 m context row and a 60 m close-up row; only the 10 m square (yellow) is judged.
  No map, NDVI or earlier label is shown. The key is kept in recheck_key.csv (do not open while labelling).
- The re-check label replaces the original label for all 30 points in the estimate (s11 variant), whatever
  direction it moves. Decoy agreement is reported as a check on labelling consistency.
Outputs: sample_recheck/{recheck_key.csv, sample_points_blind.csv, chips/R-xx.png, V3_recheck.xlsx},
         qgis/recheck_squares.gpkg
"""
import os
import numpy as np, pandas as pd, rasterio, xarray as xr, geopandas as gpd, matplotlib
from shapely.geometry import box
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import Rectangle
import s05_chips as C
from v3cfg import V3

OUT = f"{V3}/sample_recheck"
os.makedirs(f"{OUT}/chips", exist_ok=True)
d = xr.open_dataset(f"{V3}/data/esk_v3_stack_10m.nc", engine="scipy").load()
dn = d.dndvi.values


def load(f):
    k = pd.read_csv(f"{V3}/{f}/sample_key.csv")
    lab = pd.read_excel(f"{V3}/{f}/V3_labelling.xlsx", sheet_name="Labels")[["point_id", "label", "loss_nearby"]]
    return k.merge(lab, on="point_id").assign(sample=f)


pts = pd.concat([load(f) for f in ("sample", "sample_supplement", "sample_supplement2")], ignore_index=True)
pts = pts[pts.label.isin(["Loss", "No loss"])].copy()
loss_strata = np.where(pts["sample"] == "sample", pts.stratum.isin([3]), pts.stratum.isin([3, 5]))
near = pts.loss_nearby.astype(str).str.lower().str.startswith("y")
px = dn[pts.row, pts.col]
pts["reason"] = ""
pts.loc[~loss_strata & (pts.label == "Loss") & (px > -0.07), "reason"] = "omission, no NDVI drop at pixel"
pts.loc[loss_strata & (pts.label == "No loss") & near, "reason"] = "commission, loss in 30 m square"
target = pts[pts.reason != ""]
rng = np.random.default_rng(20230227)
decoy = pts[pts.reason == ""].sample(10, random_state=20230227).assign(reason="decoy")
sel = pd.concat([target, decoy]).sample(frac=1, random_state=rng.integers(1e9)).reset_index(drop=True)
sel["recheck_id"] = [f"R-{i:02d}" for i in range(1, len(sel) + 1)]
print(len(target), "target points,", len(decoy), "decoys")

key = f"{OUT}/recheck_key.csv"
if os.path.exists(f"{OUT}/V3_recheck.xlsx") and os.path.exists(key):
    raise SystemExit("re-check workbook already exists; not overwriting")
sel[["recheck_id", "point_id", "sample", "stratum", "row", "col", "easting", "northing", "label", "loss_nearby", "reason"]].to_csv(key, index=False)
sel[["recheck_id", "easting", "northing"]].to_csv(f"{OUT}/sample_points_blind.csv", index=False)

# 10 m squares for QGIS (blind ids only)
gpd.GeoDataFrame({"recheck_id": sel.recheck_id},
                 geometry=[box(x - 5, y - 5, x + 5, y + 5) for x, y in zip(sel.easting, sel.northing)],
                 crs=2193).to_file(f"{V3}/qgis/recheck_squares.gpkg", layer="recheck_squares", driver="GPKG")

# chips: row 1 = 200 m context, row 2 = 60 m close-up (same sources as s05; no post Sentinel-2, no map)
tiles = {n: C.tile_bounds(n) for n, _ in C.SOURCES if n != "s2_pre"}
s2 = np.dstack([d[f"pre_{b}"].values for b in ("red", "green", "blue")]); xs, ys = d.x.values, d.y.values
for p in sel.itertuples():
    rows = []
    for half in (100.0, 30.0):
        C.HALF = half
        rows.append((half, {n: C.read(t, p.easting, p.northing) for n, t in tiles.items()}))
    panels = [s for s in C.SOURCES if s[0] == "s2_pre" or s[0] != "post_aerial_0p1m" or rows[0][1][s[0]][0] is not None]
    fig, axs = plt.subplots(2, len(panels), figsize=(4 * len(panels), 8.6))
    for ri, (half, got) in enumerate(rows):
        ext = (p.easting - half, p.easting + half, p.northing - half, p.northing + half)
        for ax, (n, title) in zip(axs[ri], panels):
            if n == "s2_pre":
                c = (xs >= ext[0]) & (xs <= ext[1]); r = (ys >= ext[2]) & (ys <= ext[3])
                sub = np.clip(s2[np.ix_(r, c)] / 0.12, 0, 1); sub[np.isnan(sub)] = 1
                ax.imshow(sub, extent=(xs[c][0] - 5, xs[c][-1] + 5, ys[r][-1] - 5, ys[r][0] + 5), interpolation="nearest")
            elif got[n][0] is None:
                ax.set_facecolor("0.85"); ax.text(0.5, 0.5, "no coverage", ha="center", transform=ax.transAxes)
            else:
                ax.imshow(C.stretch(got[n][0]), extent=ext)
            ax.add_patch(Rectangle((p.easting - 5, p.northing - 5), 10, 10, fill=False, ec="yellow", lw=1.8 if ri == 0 else 2.5))
            if ri == 1:   # 10 m grid lines around the target in the close-up
                for k in (-15, 15):
                    ax.axvline(p.easting + k, color="white", lw=0.6, ls=(0, (3, 3)))
                    ax.axhline(p.northing + k, color="white", lw=0.6, ls=(0, (3, 3)))
            ax.set_xlim(ext[:2]); ax.set_ylim(ext[2:]); ax.set_xticks([]); ax.set_yticks([])
            ax.set_title(f"{title}{'  ·  60 m close-up' if ri else '  ·  200 m'}", fontsize=9.5)
    fig.suptitle(p.recheck_id, fontsize=13, fontweight="bold")
    fig.tight_layout(); fig.savefig(f"{OUT}/chips/{p.recheck_id}.png", dpi=110); plt.close(fig)
    print(p.recheck_id, flush=True)


def write_workbook(sel, rule):
    """Same layout as the s06 labelling sheets: chip hyperlinks and drop-down lists."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = Workbook(); ins = wb.active; ins.title = "How to label"
    for i, t in enumerate(rule, 1):
        ins.cell(i, 1, t).alignment = Alignment(wrap_text=True)
    ins["A1"].font = Font(bold=True, size=14); ins.column_dimensions["A"].width = 130
    ws = wb.create_sheet("Labels")
    ws.append(["recheck_id", "chip", "pixel_label", "confidence", "offset_seen", "notes", "easting", "northing"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="2F5D50")
    for i, r in enumerate(sel.itertuples(), 2):
        ws.cell(i, 1, r.recheck_id)
        link = ws.cell(i, 2, "open chip")
        link.hyperlink = f"chips/{r.recheck_id}.png"; link.font = Font(color="0563C1", underline="single")
        ws.cell(i, 7, r.easting); ws.cell(i, 8, r.northing)
    last = len(sel) + 1
    for col, opts in (("C", "Loss,No loss,No canopy before,Can't tell"), ("D", "High,Low"), ("E", "Yes,No")):
        dv = DataValidation(type="list", formula1=f'"{opts}"', allow_blank=True)
        ws.add_data_validation(dv); dv.add(f"{col}2:{col}{last}")
    for col, w in zip("ABCDEFGH", (11, 11, 20, 12, 12, 45, 12, 12)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"; wb.active = 1
    wb.save(f"{OUT}/V3_recheck.xlsx")

RULE = [
    "V3 pixel re-check: 30 points, about 30-40 minutes. Rule fixed before labelling.",
    "",
    "Judge ONLY the solid yellow 10 m square. Ignore everything outside it, even a slip right next to it.",
    "Use the 60 m close-up (bottom row) for the decision and the 200 m row for context.",
    "White dashed lines in the close-up mark the 30 m neighbourhood; they are for orientation only.",
    "",
    "Loss            = trees or young planted trees in the square just before (panel 2 green; aerial shows canopy)",
    "                  and, in the after image, removed, flattened or buried over about half or more of the square.",
    "No loss         = canopy still covers more than half of the square after, even if a slip passes nearby.",
    "No canopy before = the square was bare or cutover just before.",
    "Can't tell      = imagery offset, shadow, cloud or blur makes the square itself unreadable.",
    "",
    "offset_seen: write Yes if the before and after images look shifted against each other (roads, ridges).",
    "Do not open recheck_key.csv or earlier labels while labelling.",
]
write_workbook(sel, RULE)
print("done", OUT)
