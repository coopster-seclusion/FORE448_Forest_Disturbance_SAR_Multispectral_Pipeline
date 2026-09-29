"""Two versions of the baseline slide (slide 7) on top of FORE448_Esk_V3_simple_v2.pptx. Each output is a complete deck;
only slide 7 differs. Run fig_baseline_v2.py first (maps) and crop Example B from F3a (see figures/clean/example_b_*_v2.png).

A: estate at the storm coloured by stand condition, whole catchment  -> FORE448_Esk_V3_simple_v2_slide7A.pptx
C: the same numbers, small locator map + Example B aerial photo       -> FORE448_Esk_V3_simple_v2_slide7C.pptx
"""
import json
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from v3cfg import V3

PRES, CLEAN = f"{V3}/presentation", f"{V3}/figures/clean"
FONT = "Segoe UI"
INK, INK2, NOTE, PALE = RGBColor(0x0B, 0x0B, 0x0B), RGBColor(0x44, 0x43, 0x40), RGBColor(0x6E, 0x6C, 0x66), RGBColor(0xF1, 0xF3, 0xEE)
s09 = json.load(open(f"{V3}/provenance/s09_estate_condition.json"))
base = json.load(open(f"{V3}/provenance/baseline_change_areas.json"))
C, E = s09["condition_ha"], s09["estate_ha"]
TITLE = "Before the storm: 9,525 ha of plantation, a third of it young"
# colours match fig_baseline_v2.py
ROWS = [(C["mature"], "mature", "1B8A5A", "0B4A2F", "closed canopy, not harvested since 2017"),
        (C["young"], "young stands", "B5D93B", "4F6A0A", "replanted or regrowing after 2017–22 harvests"),
        (C["open"], "open cutover", "D9C7A0", "8A7650", "bare at the storm: no canopy to lose")]
HOW = [("How we built it: ", {"bold": True, "color": INK}),
       ("a satellite classifier (AlphaEarth 2022) plus Forestry Catchment Planner stands, then each 10 m pixel sorted by "
        "Sentinel-2 greenness just before the storm.", {})]
WHY = [("Why it matters: ", {"bold": True, "color": INK}),
       ("young stands lost 15% of their canopy against 4% for mature. The 2018/19 national map lists "
        f"{base['estate_in_lcdb6_harvested']:,} ha of them as 'harvested', so a map of standing forest would miss them.", {})]
FOOT = (f"Estate agrees with LCDB v6.0 2018/19 exotic + harvested forest on {100 * base['kept'] / E:.0f}% of its area. "
        f"{C['no_optical']} ha had no cloud-free data. At 10 m the estate edges bleed into roads and gully margins; the 466 ha "
        "estimate is unaffected because the checkers judged every dot.")


def _run(p, t, size, color, bold=False):
    r = p.add_run(); r.text = t; r.font.name = FONT; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color


def text(s, x, y, w, h, t, size=18, color=INK2, bold=False, anchor=MSO_ANCHOR.TOP):
    tf = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    p = tf.paragraphs[0]
    for seg, st in ([(t, {})] if isinstance(t, str) else t):
        _run(p, seg, st.get("size", size), st.get("color", color), st.get("bold", bold))
    return tf


def rect(s, x, y, w, h, fill, edge=None):
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    r.fill.solid(); r.fill.fore_color.rgb = fill; r.shadow.inherit = False
    if edge is None:
        r.line.fill.background()
    else:
        r.line.color.rgb = edge; r.line.width = Pt(1.5)
    return r


def picture(s, path, x, y, w, h):
    im = Image.open(path); ar = im.width / im.height
    ww, hh = (w, w / ar) if w / ar <= h else (h * ar, h)
    return s.shapes.add_picture(path, Inches(x + (w - ww) / 2), Inches(y + (h - hh) / 2), Inches(ww), Inches(hh))


def rows(s, x, y, w, step, head_size, sub_size):
    for k, (ha, name, fill, edge, sub) in enumerate(ROWS):
        yy = y + k * step
        rect(s, x, yy + 0.04, 0.32, 0.32, RGBColor.from_string(fill), RGBColor.from_string(edge))
        text(s, x + 0.5, yy, w - 0.5, 0.4, f"{ha:,} ha {name} ({100 * ha / E:.0f}%)", size=head_size, color=INK, bold=True)
        text(s, x + 0.5, yy + 0.38, w - 0.5, 0.4, sub, size=sub_size)


def start(title=TITLE):
    prs = Presentation(f"{PRES}/FORE448_Esk_V3_simple_v2.pptx")
    s = prs.slides[6]
    assert s.shapes[0].text_frame.text.startswith("We rebuilt the plantation map"), "slide 7 is not the baseline slide"
    for sh in list(s.shapes)[1:]:
        sh._element.getparent().remove(sh._element)
    p = s.shapes[0].text_frame.paragraphs[0]
    for r in list(p.runs)[1:]:
        r._r.getparent().remove(r._r)
    p.runs[0].text = title
    return prs, s


NOTES_CORE = (f"Before measuring loss we needed the plantation as it stood just before the storm: {E:,} ha, about a third of "
              "the catchment. We built it from a satellite classifier plus Forestry Catchment Planner stands, then sorted "
              f"every 10 m pixel by how green it was just before the storm: {C['mature']:,} ha mature, {C['young']:,} ha young "
              f"stands, {C['open']} ha open cutover. ")
NOTES_WHY = ("A third of it was young, replanted after the 2017 to 2022 harvests, and that matters because young stands "
             "were hit hardest. The 2018/19 national map still lists much of this as 'harvested', so a map of standing forest "
             "would have missed it.")

# ---------- A: whole-catchment map ----------
prs, s = start()
text(s, 0.6, 1.35, 7.7, 0.9, HOW, size=16)
rows(s, 0.6, 2.55, 7.7, 0.85, 20, 15)
rect(s, 0.6, 5.2, 7.7, 1.0, PALE)
text(s, 0.85, 5.2, 7.25, 1.0, WHY, size=15, anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.6, 6.4, 7.7, 0.9, FOOT, size=12, color=NOTE, anchor=MSO_ANCHOR.BOTTOM)
picture(s, f"{CLEAN}/estate_age_map_v2.png", 8.7, 1.15, 4.1, 6.2)
s.notes_slide.notes_text_frame.text = "[Presenter 2] ~35 s. " + NOTES_CORE + "On the map: dark green mature, lime young, tan cutover. " + NOTES_WHY
prs.save(f"{PRES}/FORE448_Esk_V3_simple_v2_slide7A.pptx"); print("saved A")

# ---------- C: locator + Example B photo ----------
prs, s = start()
text(s, 0.6, 1.35, 4.4, 1.3, HOW, size=15)
rows(s, 0.6, 2.7, 4.4, 0.83, 17, 13)
rect(s, 0.6, 5.2, 4.4, 1.25, PALE)
text(s, 0.8, 5.2, 4.0, 1.25, WHY, size=13, anchor=MSO_ANCHOR.MIDDLE)
picture(s, f"{CLEAN}/estate_age_locator_v2.png", 5.2, 1.35, 2.3, 4.6)
picture(s, f"{CLEAN}/example_b_chip_v2.png", 7.75, 1.3, 5.0, 5.0)
text(s, 7.75, 6.35, 5.0, 0.6, [("Example B, 2021–22 aerial: ", {"bold": True, "color": INK}),
                               ("rows of young pines inside a polygon the 2018/19 national map calls 'harvested' (yellow).", {})],
     size=13)
text(s, 0.6, 6.6, 7.0, 0.75, FOOT, size=11, color=NOTE, anchor=MSO_ANCHOR.BOTTOM)
s.notes_slide.notes_text_frame.text = ("[Presenter 2] ~40 s. " + NOTES_CORE + "Here's what 'young' looks like on the ground: "
                                       "Example B, on the 2021–22 aerial photo, is rows of young pines. The national map still "
                                       "has it as 'harvested', the yellow outline. " + NOTES_WHY)
prs.save(f"{PRES}/FORE448_Esk_V3_simple_v2_slide7C.pptx"); print("saved C")
