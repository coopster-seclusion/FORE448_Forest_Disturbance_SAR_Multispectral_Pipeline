"""V3 step 6: blind labelling workbook (sample/V3_labelling.xlsx) with dropdowns and chip links."""
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

import os, sys
from v3cfg import RES, SAMPLE
if os.path.exists(f"{SAMPLE}/V3_labelling.xlsx"):
    sys.exit(f"{SAMPLE}/V3_labelling.xlsx exists (may hold labels); delete it deliberately to rebuild")
SUPPLEMENT = "supplement" in SAMPLE
YOUNG = SAMPLE.rstrip("/").endswith("supplement2")  # young-stand draw: clearer rule + before_cover
pts = pd.read_csv(f"{SAMPLE}/sample_points_blind.csv")

wb = Workbook()
ins = wb.active
ins.title = "How to label"
guide = [
    f"V3 reference labelling: {len(pts)} points" + (" (supplement: young stands and newly added plantation)" if SUPPLEMENT else ""),
    "",
    f"For each row, click the chip link. Look at the YELLOW {RES} m square in every panel (the dashed square is {3 * RES} m).",
    "Panels: (1) aerial 2021-22, (2) Sentinel-2 Jan-Feb 2023 = just before the cyclone, (3) satellite 21 Feb 2023 = just after, (4) 0.1 m aerial Feb 2023 where it exists.",
    "",
    "Choose ONE label for the yellow square:",
    "  Loss - tree canopy just before (panel 2) and gone or damaged over about half or more of the square after (slip scar, bare soil, flattened or missing trees, silt).",
    "  No loss - tree canopy before and still intact after.",
    "  No canopy before - already open, cutover or very young trees just before the cyclone (look at panel 2 if panels 1 and 2 disagree: felled between 2021 and Jan 2023 counts here).",
    "  Can't tell - cloud, deep shadow, blur or no image.",
    *(["  Not plantation - the square is native bush/scrub, pasture or other land (not planted forest)."] if SUPPLEMENT else []),
    "",
    "",
    *(["",
       "YOUNG-STAND RULE for these points: most are young plantation or recently harvested land.",
       "  Loss = young trees or regrowth vegetation that was there just before (panel 2 green) is removed or buried by a slip, debris or silt.",
       "  No canopy before = the square was bare cutover, slash or soil just before the cyclone (panel 2 brown/bare).",
       "  before_cover (optional but useful): what was in the square just before - Young planted trees / Weedy cutover or regrowth / Mature trees / Bare.",
       ] if YOUNG else []),
    "Then answer loss_nearby: is there any cyclone canopy loss inside the DASHED square (Yes/No)?",
    "  This records near-misses caused by image offsets and mixed pixels. The label column above stays strict to the yellow square.",
    "",
    "Optional: Cause (for Loss), Confidence (High/Low), and Notes.",
    "Do not look at the V3 loss map while labelling. Judge the square, not the surroundings.",
    "Save the file when finished. Nothing else is needed.",
]
for i, t in enumerate(guide, 1):
    ins.cell(i, 1, t).alignment = Alignment(wrap_text=True)
ins["A1"].font = Font(bold=True, size=14)
ins.column_dimensions["A"].width = 130

ws = wb.create_sheet("Labels")
head = ["point_id", "chip", "label", "loss_nearby", "cause", "confidence", "notes", "easting", "northing"] + (["before_cover"] if YOUNG else [])
ws.append(head)
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="2F5D50")
for i, r in pts.iterrows():
    row = i + 2
    ws.cell(row, 1, r.point_id)
    link = ws.cell(row, 2, "open chip")
    link.hyperlink = f"chips/{r.point_id}.png"; link.font = Font(color="0563C1", underline="single")
    ws.cell(row, 8, r.easting); ws.cell(row, 9, r.northing)
last = len(pts) + 1
for col, opts in (("C", "Loss,No loss,No canopy before,Can't tell" + (",Not plantation" if SUPPLEMENT else "")),
                  ("D", "Yes,No"),
                  ("E", "Slip or debris flow,Windthrow,Flood or silt,Other"),
                  ("F", "High,Low"), *([("J", "Young planted trees,Weedy cutover or regrowth,Mature trees,Bare")] if YOUNG else [])):
    dv = DataValidation(type="list", formula1=f'"{opts}"', allow_blank=True)
    ws.add_data_validation(dv); dv.add(f"{col}2:{col}{last}")
for col, w in zip("ABCDEFGHIJ", (10, 11, 18, 12, 20, 12, 40, 12, 12, 26)):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "A2"
wb.active = 1
wb.save(f"{SAMPLE}/V3_labelling.xlsx")
print("wrote", last - 1, "rows")
