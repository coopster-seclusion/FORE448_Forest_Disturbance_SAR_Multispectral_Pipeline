"""Add the location map and the all-loss-areas map to the finished simple deck in place (keeps hand edits).

1. Slide "The Esk: steep hill country..." : its picture is replaced by figures/clean/location_map.png.
2. New slide "Where canopy was lost across the catchment" (bullets + figures/clean/loss_areas_map.png), inserted after
   the headline-result slide ("About 470 ha of plantation canopy was lost"). Skipped if it already exists.
"""
from PIL import Image
from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.ns import qn
from v3cfg import V3

PATH = f"{V3}/presentation/FORE448_Esk_V3_simple.pptx"
FONT, INK, INK2, MUTED = "Segoe UI", RGBColor(0x0B, 0x0B, 0x0B), RGBColor(0x44, 0x43, 0x40), RGBColor(0x8A, 0x88, 0x80)
prs = Presentation(PATH)


def slide_titled(start):
    for i, s in enumerate(prs.slides):
        if any(sh.has_text_frame and sh.text_frame.text.strip().startswith(start) for sh in s.shapes):
            return i, s
    return None, None


def tf_at(s, x, y, w, h):
    tf = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame
    tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.TOP
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    return tf


def run(p, t, size, color, bold=False):
    r = p.add_run(); r.text = t; r.font.name = FONT; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color


def fit_picture(s, path, x, y, w, h):
    im = Image.open(path); ar = im.width / im.height
    ww, hh = (w, w / ar) if w / ar <= h else (h * ar, h)
    s.shapes.add_picture(path, Inches(x + (w - ww) / 2), Inches(y + (h - hh) / 2), Inches(ww), Inches(hh))


# 1. location map on the catchment slide
i, s = slide_titled("The Esk: steep hill country")
pic = max((sh for sh in s.shapes if sh.shape_type == 13), key=lambda sh: sh.width * sh.height)
pic._element.getparent().remove(pic._element)
fit_picture(s, f"{V3}/figures/clean/location_map.png", 5.8, 1.45, 7.0, 5.45)
print("slide", i + 1, "← location_map.png")

# 2. all loss areas slide after the headline result
if slide_titled("Where canopy was lost across the catchment")[1] is None:
    hi, _ = slide_titled("About 470 ha of plantation canopy was lost")
    ns = prs.slides.add_slide(next(l for l in prs.slide_layouts if l.name == "Empty"))
    tf = tf_at(ns, 0.6, 0.42, 12.1, 0.9); run(tf.paragraphs[0], "Where canopy was lost across the catchment", 30, INK, True)
    tf = tf_at(ns, 0.6, 1.75, 5.6, 5.0)
    items = [("Orange: ", "823 ha of plantation canopy mapped as lost"), ("Blue: ", "about 650 ha of native canopy mapped as lost"),
             "Loss clusters in gullies, along streams and on the valley floor",
             "The map shows where; the checked estimate (466 ha) shows how much"]
    for k, it in enumerate(items):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        pPr = p._p.get_or_add_pPr(); pPr.set("marL", str(Inches(0.3))); pPr.set("indent", str(-Inches(0.3)))
        etree.SubElement(pPr, qn("a:buFont")).set("typeface", "Arial"); etree.SubElement(pPr, qn("a:buChar")).set("char", "•")
        p.space_after = Pt(14); p.line_spacing = 1.1
        if isinstance(it, tuple):
            run(p, it[0], 20, INK, True); run(p, it[1], 20, INK2)
        else:
            run(p, it, 20, INK2)
    fit_picture(ns, f"{V3}/figures/clean/loss_areas_map.png", 6.4, 1.3, 6.4, 5.65)
    tf = tf_at(ns, 0.6, 7.0, 5.6, 0.3); run(tf.paragraphs[0], "Map-based, 10 m Sentinel-2 (before sample correction).", 11, MUTED)
    ns.notes_slide.notes_text_frame.text = (
        "[Presenter 3] ~30 s. This is every area the 10 m map flagged. Orange is plantation canopy lost, 823 ha on the map; "
        "blue is native canopy lost, about 650 ha. Loss clusters in gullies, along streams and on the valley floor. The map "
        "shows where; the checked estimate of 466 ha says how much, because the map alone overstated mature loss.")
    ids = prs.slides._sldIdLst; new = ids[-1]; ids.remove(new); ids.insert(hi + 1, new)
    print("new slide", hi + 2, "← loss_areas_map.png")
prs.save(PATH)
